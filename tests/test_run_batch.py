"""Tests for run_batch() and plan_batch(): one new Excel per run, patient IDs,
the work decided for each file, transcript reuse, progress and errors."""

import json
from datetime import datetime

import pytest
from openpyxl import load_workbook

from phenoscribe import pipeline
from phenoscribe.config import Config, PatientConfig
from phenoscribe.output import write_excel
from phenoscribe.pipeline import (
    batch_patient_id,
    new_results_path,
    plan_batch,
    run_batch,
    transcript_cache_path,
)

CONFIG = Config(patient=PatientConfig(id_prefix="MGA."))


def _files(folder, *names, content=b"audio"):
    paths = []
    for name in names:
        path = folder / name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(content)
        paths.append(str(path))
    return paths


@pytest.fixture
def fake_pipeline(monkeypatch):
    """Replace process_recording: one Fatigue code per file, or raise for 'bad' files."""
    calls = []

    def fake(input_path, patient_id, config, output_path=None, on_step=None, **kwargs):
        calls.append((patient_id, kwargs.get("reuse_transcript")))
        on_step("removing personal details")
        if "bad" in input_path:
            raise ValueError("Unsupported file type '.xyz'")
        matches = [{"hpo_id": "HP:0012378", "hpo_term": "Fatigue", "patient_verbatim": ""}]
        write_excel(patient_id, matches, output_path, fmt="semicolon")
        return matches

    monkeypatch.setattr(pipeline, "process_recording", fake)
    return calls


def _patient_ids(path):
    ws = load_workbook(path).worksheets[0]
    return [row[0] for row in ws.iter_rows(min_row=2, values_only=True)]


def test_two_runs_give_two_files(tmp_path, fake_pipeline):
    files = _files(tmp_path / "in", "467.mp3")
    first = run_batch(files, CONFIG, tmp_path)
    second = run_batch(files, CONFIG, tmp_path)

    assert first.excel_path != second.excel_path
    assert first.excel_path.name.startswith("results_")
    # The rerun adds no rows to the first file.
    assert _patient_ids(first.excel_path) == ["MGA.467"]
    assert _patient_ids(second.excel_path) == ["MGA.467"]


def test_digit_stem_gives_mga_id():
    assert batch_patient_id("467.mp3", "MGA.") == "MGA.467"


def test_name_in_filename_gives_hash_id():
    pid = batch_patient_id("jean-martin.mp3", "MGA.")
    assert pid.startswith("pt-") and len(pid) == len("pt-") + 8
    assert "jean" not in pid
    assert batch_patient_id("jean-martin.mp3", "MGA.") == pid  # stable
    # A recording and its transcript share one ID.
    assert batch_patient_id("jean-martin.txt", "MGA.") == pid


def test_results_path_is_dated_and_never_reused(tmp_path):
    now = datetime(2026, 10, 22, 14, 5)
    first = new_results_path(tmp_path, now)
    assert first.name == "results_2026-10-22_1405.xlsx"
    first.touch()
    assert new_results_path(tmp_path, now).name == "results_2026-10-22_1405_2.xlsx"


def test_a_failed_file_does_not_stop_the_batch(tmp_path, fake_pipeline):
    files = _files(tmp_path / "in", "bad.mp3", "467.mp3")
    batch = run_batch(files, CONFIG, tmp_path)

    bad, good = batch.files
    assert isinstance(bad.error, ValueError)
    assert good.error is None and good.n_codes == 1
    assert good.matches[0]["hpo_id"] == "HP:0012378"
    assert _patient_ids(batch.excel_path) == ["MGA.467"]


def test_no_excel_when_every_file_fails(tmp_path, fake_pipeline):
    files = _files(tmp_path / "in", "bad.mp3")
    assert run_batch(files, CONFIG, tmp_path).excel_path is None


def test_progress_names_the_file_and_the_step(tmp_path, fake_pipeline):
    messages = []
    run_batch(
        _files(tmp_path / "in", "466.mp3", "467.mp3"),
        CONFIG,
        tmp_path,
        on_progress=lambda fraction, message: messages.append((fraction, message)),
    )
    # File 2 of 2, past transcription: (1 + 0.6) / 2.
    assert (0.8, "2/2 · 467.mp3 · removing personal details") in messages
    assert messages[-1] == (1.0, "Finished")


def test_filename_mapping_keeps_the_name_out_of_the_excel(tmp_path, fake_pipeline):
    batch = run_batch(_files(tmp_path / "in", "jean-martin.mp3"), CONFIG, tmp_path)
    pid = batch.files[0].patient_id

    mapping = json.loads((tmp_path / "filename_mapping.json").read_text(encoding="utf-8"))
    assert mapping == {pid: "jean-martin.mp3"}
    assert _patient_ids(batch.excel_path) == [pid]


# --- plan_batch: the work for each file ----------------------------------------


def _actions(plans):
    return {p.name: (p.action, p.note) for p in plans}


def test_plan_audio_transcript_and_unsupported(tmp_path):
    files = _files(tmp_path / "in", "101.mp3", "103.txt", "consent.pdf", ".DS_Store", "103_pseudo.txt")
    assert _actions(plan_batch(files, CONFIG, tmp_path)) == {
        "101.mp3": ("transcribe", ""),
        "103.txt": ("transcript", ""),
        "consent.pdf": ("skip", "Not an audio file or transcript"),
    }


def test_plan_transcript_wins_over_its_audio(tmp_path):
    files = _files(tmp_path / "in", "467.mp3", "467.txt")
    assert _actions(plan_batch(files, CONFIG, tmp_path)) == {
        "467.mp3": ("skip", "Uses 467.txt instead"),
        "467.txt": ("transcript", ""),
    }


def test_plan_skips_two_files_for_one_patient(tmp_path):
    # 467.mp3 in two folders would overwrite each other's transcript.
    files = _files(tmp_path / "a", "467.mp3") + _files(tmp_path / "b", "467.mp3")
    plans = plan_batch(files, CONFIG, tmp_path)
    assert [p.action for p in plans] == ["skip", "skip"]
    assert plans[0].note == "Two files give patient MGA.467. Rename one."


def test_plan_reuses_the_transcript_of_the_same_audio(tmp_path):
    (audio,) = _files(tmp_path / "in", "467.mp3", content=b"bytes one")
    cache = transcript_cache_path(audio, CONFIG, str(tmp_path / "transcripts"))
    cache.parent.mkdir(parents=True)
    cache.write_text("bonjour", encoding="utf-8")
    assert plan_batch([audio], CONFIG, tmp_path)[0].action == "reuse"

    # Other bytes under the same name: transcribe again.
    (tmp_path / "in" / "467.mp3").write_bytes(b"bytes two")
    assert plan_batch([audio], CONFIG, tmp_path)[0].action == "transcribe"


def test_cache_key_changes_with_the_whisper_settings(tmp_path):
    (audio,) = _files(tmp_path / "in", "467.mp3")
    other = Config()
    other.transcription.language = "nl"
    assert transcript_cache_path(audio, CONFIG, "t") != transcript_cache_path(audio, other, "t")


def test_run_batch_skips_planned_skips_and_honours_retranscribe(tmp_path, fake_pipeline):
    files = _files(tmp_path / "in", "466.mp3", "467.mp3", "467.txt", "notes.pdf")
    run_batch(files, CONFIG, tmp_path, retranscribe=frozenset({"466.mp3"}))
    assert fake_pipeline == [("MGA.466", False), ("MGA.467", True)]


def test_process_recording_reuses_and_saves_the_audio_transcript(tmp_path, monkeypatch):
    """The real process_recording, with every step after transcription stubbed."""
    (audio,) = _files(tmp_path / "in", "467.mp3")
    monkeypatch.setattr(pipeline, "check_obo_version", lambda *a: "hp/releases/2026-02-16")
    monkeypatch.setattr(pipeline, "find_personal_details", lambda text, **k: [])
    monkeypatch.setattr(pipeline, "extract_symptoms", lambda text, **k: [])
    monkeypatch.setattr(pipeline, "match_hpo", lambda symptoms, **k: [])
    monkeypatch.setattr(pipeline, "write_excel", lambda *a, **k: None)
    whisper = []
    monkeypatch.setattr(pipeline, "transcribe", lambda path, **k: whisper.append(path) or "bonjour")
    out = str(tmp_path / "out" / "results.xlsx")

    for reuse in (True, True, False):
        pipeline.process_recording(
            audio, "MGA.467", CONFIG, output_path=out, reuse_transcript=reuse, no_review=True
        )

    assert len(whisper) == 2  # the second run reused the saved transcript
    cache = transcript_cache_path(audio, CONFIG, str(tmp_path / "out" / "transcripts"))
    assert cache.read_text(encoding="utf-8") == "bonjour"


def test_the_bar_moves_within_the_matching_step():
    from phenoscribe.pipeline import CODE_PROGRESS, step_fraction

    assert step_fraction(CODE_PROGRESS, "matching HPO codes") == 0.4
    assert step_fraction(CODE_PROGRESS, "matching HPO codes · symptom 1 of 10") == 0.4
    assert step_fraction(CODE_PROGRESS, "matching HPO codes · symptom 6 of 10") == pytest.approx(0.65)
    assert step_fraction(CODE_PROGRESS, "writing the Excel file") == 0.9


@pytest.mark.parametrize("name, patient, kind, date", [
    ("MGA.363 interview.docx", "MGA.363", "interview", ""),
    ("363  interview evolution  2026-10-15_11-09-09.docx", "MGA.363", "follow-up", "2026-10-15"),
    ("MGA.363 mail du 20260930.docx", "MGA.363", "email", "2026-09-30"),
    ("006 interview evolution  20260922 sept. à 10-45.docx", "MGA.006", "follow-up", "2026-09-22"),
    ("100 intrview evolution 20260922.docx", "MGA.100", "follow-up", "2026-09-22"),
    ("XX_ interview 20230827_.docx", "", "interview", "2023-08-27"),
    ("467.mp3", "MGA.467", "interview", ""),
])
def test_interview_details_link_the_files_of_one_patient(name, patient, kind, date):
    from phenoscribe.pipeline import interview_details

    assert interview_details(name, "MGA.") == {"patient": patient, "interview": kind, "date": date}


def test_the_detailed_excel_has_patient_interview_and_date(tmp_path):
    path = tmp_path / "r.xlsx"
    matches = [{"hpo_id": "HP:0012378", "hpo_term": "Fatigue", "patient_verbatim": "épuisée"}]
    write_excel("pt-1a2b3c4d", matches, str(path), fmt="detailed",
                interview={"patient": "MGA.363", "interview": "follow-up", "date": "2026-10-15"})
    ws = load_workbook(path).worksheets[0]
    row = dict(zip([c.value for c in ws[1]], [c.value for c in ws[2]]))
    assert (row["Patient_ID"], row["Patient"], row["Interview"], row["Date"]) == (
        "pt-1a2b3c4d", "MGA.363", "follow-up", "2026-10-15")


def test_rows_added_to_an_older_workbook_get_the_new_headers(tmp_path):
    from openpyxl import Workbook

    from phenoscribe.output import HEADERS

    path = tmp_path / "old.xlsx"
    wb = Workbook()
    wb.active.append(HEADERS["detailed"][:14])  # the layout before 2026-10-27
    wb.save(path)
    write_excel("MGA.363", [{"hpo_id": "HP:0012378", "hpo_term": "Fatigue", "patient_verbatim": ""}],
                str(path), fmt="detailed", interview={"patient": "MGA.363", "interview": "interview", "date": ""})
    ws = load_workbook(path).worksheets[0]
    assert [c.value for c in ws[1]] == HEADERS["detailed"]
    assert ws.cell(row=2, column=15).value == "MGA.363"
