"""Who brings up a finding: the patient, the doctor with the patient's yes, or the doctor only.

A finding that only the doctor says is not a finding of the patient. It stays
in the detailed sheet as "Doctor only" and is not counted.
"""

import openpyxl

from phenoscribe.aggregate import load_patient_codes
from phenoscribe.extract_symptoms import _parse_response
from phenoscribe.match_hpo import finding_status, present_codes
from phenoscribe.output import dedup_by_code, write_excel


def _match(code, source, negated=False, term="Fatigue"):
    return {"hpo_id": code, "hpo_term": term, "clinical_term": term.lower(),
            "patient_verbatim": "", "negated": negated, "source": source}


def test_parser_reads_source():
    found = _parse_response("""[
      {"clinical_term": "fatigue", "source": "patient"},
      {"clinical_term": "palpitations", "source": "Confirmed"},
      {"clinical_term": "anosmia", "source": "doctor"},
      {"clinical_term": "fever", "source": "nurse"},
      {"clinical_term": "cough"}
    ]""")
    assert [f["source"] for f in found] == ["patient", "confirmed", "doctor", "", ""]


def test_status_and_present_codes():
    patient, doctor = _match("HP:1", "patient"), _match("HP:2", "doctor")
    denied = _match("HP:3", "doctor", negated=True)
    old = _match("HP:4", "")  # a finding from before the source field
    assert [finding_status(m) for m in (patient, doctor, denied, old)] == [
        "Present", "Doctor only", "Absent", "Present"]
    assert present_codes([patient, doctor, denied, old]) == [patient, old]


def test_restated_finding_keeps_the_strongest_source():
    rows = dedup_by_code([_match("HP:1", "doctor"), _match("HP:1", "confirmed")], "detailed")
    assert [r["source"] for r in rows] == ["confirmed"]


def test_doctor_only_in_the_sheets(tmp_path):
    matches = [_match("HP:0012378", "patient"),
               _match("HP:0001945", "doctor", term="Fever")]
    detailed, semicolon = tmp_path / "d.xlsx", tmp_path / "s.xlsx"
    write_excel("MGA.001", matches, str(detailed), fmt="detailed")
    write_excel("MGA.001", matches, str(semicolon), fmt="semicolon")

    rows = list(openpyxl.load_workbook(detailed).active.iter_rows(values_only=True))
    head = rows[0]
    by_code = {r[2]: (r[head.index("Present/Absent")], r[head.index("Source")]) for r in rows[1:]}
    assert by_code == {"HP:0012378": ("Present", "patient"),
                       "HP:0001945": ("Doctor only", "doctor")}

    row = list(openpyxl.load_workbook(semicolon).active.iter_rows(values_only=True))[1]
    assert "HP:0012378" in row[1] and "HP:0001945" not in (row[1] or "") + (row[2] or "")

    assert [e["hpo_id"] for e in load_patient_codes(str(detailed))["MGA.001"]] == ["HP:0012378"]


def test_a_stronger_source_replaces_the_doctor_row():
    doctor = dict(_match("HP:1", "doctor"), patient_verbatim="vous aviez des nausées", onset="adult")
    patient = dict(_match("HP:1", "patient"), patient_verbatim="j'ai la nausée")
    rows = dedup_by_code([doctor, patient], "detailed")
    assert [(r["source"], r["patient_verbatim"], r.get("onset")) for r in rows] == [
        ("patient", "j'ai la nausée", None)]
    # A doctor row after a patient row adds no words.
    rows = dedup_by_code([patient, doctor], "detailed")
    assert [r["patient_verbatim"] for r in rows] == ["j'ai la nausée"]


def test_mentions_of_one_finding_share_its_status():
    from phenoscribe.timeline import mention_rows
    doctor = dict(_match("HP:1", "doctor"), mentions=["vous aviez des nausées"])
    patient = dict(_match("HP:1", "confirmed"), mentions=["oui la nausée"])
    rows = mention_rows([doctor, patient], "Vous aviez des nausées ? Oui la nausée.")
    assert {r["status"] for r in rows} == {"Present"}


def test_an_empty_status_cell_still_counts(tmp_path):
    path = tmp_path / "hand.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["Patient_ID", "HPO Term", "HPO Code", "Present/Absent"])
    wb.active.append(["MGA.001", "Fatigue", "HP:0012378", None])
    wb.active.append(["MGA.001", "Fever", "HP:0001945", "Doctor only"])
    wb.save(path)
    assert [e["hpo_id"] for e in load_patient_codes(str(path))["MGA.001"]] == ["HP:0012378"]
