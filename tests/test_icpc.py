"""Tests for the ICPC-3 / ICPC-2 columns next to the HPO code.

Marc asked for ICPC codes so he can count the corpus with the classification
that family medicine uses. The table covers the frequent terms only; an HPO
code outside it must give empty cells, never a guess.
"""

import csv
import importlib.util
import re
import subprocess
import sys
import zipfile
from pathlib import Path

import openpyxl
import pytest
from openpyxl import Workbook

from phenoscribe.icpc import chapter_and_component, icpc_for, load_icpc_table
from phenoscribe.output import HEADERS, write_excel

ROOT = Path(__file__).resolve().parent.parent
TABLE = ROOT / "src" / "phenoscribe" / "resources" / "hpo_icpc.csv"


def _load_script(name):
    # scripts/ is not a package; load the module by path.
    spec = importlib.util.spec_from_file_location(name, ROOT / "scripts" / f"{name}.py")
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


build_icpc_table = _load_script("build_icpc_table")
add_icpc_columns = _load_script("add_icpc_columns")


def test_committed_table_is_valid():
    with TABLE.open(newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    assert len(rows) > 300
    ids = [r["hpo_id"] for r in rows]
    assert len(ids) == len(set(ids))
    for r in rows:
        assert re.fullmatch(r"HP:\d{7}", r["hpo_id"])
        assert re.fullmatch(r"(\d[A-Z]|[A-Z]{2})\d{2}(\.\d{2})?", r["icpc3"]), r
        assert r["icpc2"] == "" or re.fullmatch(r"[A-Z]\d{2}", r["icpc2"]), r
        assert r["source"] in ("marc", "draft")


FATIGUE = {
    "chapter": "A General",
    "component": "Symptoms, complaints and abnormal findings",
    "icpc3": "AS04",
    "icpc2": "A04",
    "source": "marc",
}
ICPC_HEADERS = ("ICPC-3 Chapter", "ICPC-3 Component", "ICPC-3", "ICPC-2", "ICPC Source")


def test_fatigue_maps_to_general_tiredness():
    assert icpc_for("HP:0012378") == FATIGUE


def test_chapter_and_component():
    assert chapter_and_component("TS99") == ("T Endocrine, metabolic and nutritional system",
                                            "Symptoms, complaints and abnormal findings")
    assert chapter_and_component("KD66") == ("K Circulatory system", "Diagnoses and diseases")
    assert chapter_and_component("ZC16") == ("Z Social problems", "Social problems")
    assert chapter_and_component("2F85") == ("II Functioning", "Functioning")
    assert chapter_and_component("") == ("", "")


def test_every_table_code_has_a_chapter():
    for row in load_icpc_table().values():
        assert row["chapter"] and row["component"], row


def test_lab_results_keep_their_body_system():
    # Hypocalcemia is a metabolic finding: the general code AS50 put it in chapter A.
    assert icpc_for("HP:0002901")["icpc3"] == "TS99"
    assert not [r for r in load_icpc_table().values() if r["icpc3"] == "AS50"]


def test_unknown_or_empty_code_gives_empty_cells():
    empty = dict.fromkeys(FATIGUE, "")
    assert icpc_for("HP:9999999") == empty
    assert icpc_for("") == empty
    assert icpc_for(None) == empty


def test_missing_table_gives_empty_cells(tmp_path):
    path = str(tmp_path / "missing.csv")
    assert load_icpc_table(path) == {}
    assert icpc_for("HP:0012378", path)["icpc3"] == ""


def test_detailed_sheet_has_icpc_columns(tmp_path):
    path = tmp_path / "out.xlsx"
    matches = [
        {"hpo_id": "HP:0012378", "hpo_term": "Fatigue", "patient_verbatim": "fatigué"},
        {"hpo_id": "HP:9999999", "hpo_term": "Unknown", "patient_verbatim": ""},
    ]
    write_excel("MGA.001", matches, str(path), fmt="detailed", hpo_release="hp/releases/2026-02-16")
    wb = openpyxl.load_workbook(path)
    rows = list(wb.active.iter_rows(values_only=True))
    first = rows[0].index(ICPC_HEADERS[0])
    assert rows[0][first:first + 5] == ICPC_HEADERS
    by_code = {r[2]: r[first:first + 5] for r in rows[1:]}
    assert by_code["HP:0012378"] == tuple(FATIGUE.values())
    assert by_code["HP:9999999"] == (None,) * 5
    prov = {r[0]: r[1] for r in wb["Provenance"].iter_rows(values_only=True)}
    assert prov["hpo_release"] == "hp/releases/2026-02-16"
    assert prov["icpc_table"].endswith("draft)")


def test_older_workbook_gets_icpc_headers_at_the_end(tmp_path):
    path = tmp_path / "old.xlsx"
    wb = Workbook()
    wb.active.append(HEADERS["detailed"][:17])  # the layout before the ICPC columns
    wb.save(path)
    write_excel("MGA.002", [{"hpo_id": "HP:0012378", "hpo_term": "Fatigue", "patient_verbatim": ""}],
                str(path), fmt="detailed")
    ws = openpyxl.load_workbook(path).active
    assert [c.value for c in ws[1]] == HEADERS["detailed"]
    first = HEADERS["detailed"].index("ICPC-3 Chapter")
    assert [c.value for c in ws[2]][first:first + 5] == list(FATIGUE.values())


def test_add_icpc_columns_keeps_other_cells(tmp_path):
    src, dst = tmp_path / "in.xlsx", tmp_path / "out.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Codes"
    ws.append(["Patient_ID", "HPO Term", "HPO Code", "Code Source"])
    ws.append(["MGA.001", "Fatigue", "HP:0012378", "hand"])
    ws.append(["MGA.001", "Uncoded", None, ""])
    wb.save(src)

    counts = add_icpc_columns.add_icpc_columns(str(src), str(dst))

    rows = list(openpyxl.load_workbook(dst)["Codes"].iter_rows(values_only=True))
    assert rows[0] == ("Patient_ID", "HPO Term", "HPO Code", "Code Source", *ICPC_HEADERS)
    assert rows[1] == ("MGA.001", "Fatigue", "HP:0012378", "hand", *FATIGUE.values())
    assert rows[2][:4] == ("MGA.001", "Uncoded", None, None)
    assert counts["marc"] == 1
    assert list(openpyxl.load_workbook(src)["Codes"].iter_rows(values_only=True))[0][-1] == "Code Source"


OBO = """format-version: 1.2

[Term]
id: HP:0012378
name: Fatigue

[Term]
id: HP:0000458
name: Anosmia

[Term]
id: HP:0002027
name: Abdominal pain

[Term]
id: HP:0001945
name: Fever
"""

BROWSER = {
    "AS04": {"found": True, "code": "AS04", "title": "t", "icpc2": ["A04"]},
    "AS03": {"found": True, "code": "AS03", "title": "t", "icpc2": ["A03"]},
    "NS08": {"found": True, "code": "NS08", "title": "t", "icpc2": ["N16"]},
    "NS11": {"found": True, "code": "NS11", "title": "t", "icpc2": ["N19", "P10"]},
    "DS01": {"found": True, "code": "DS01", "title": "t", "icpc2": ["D01"]},
}


def test_build_applies_fixes_drafts_and_sources(tmp_path):
    obo = tmp_path / "hp.obo"
    obo.write_text(OBO)
    marc = {"HP:0012378": "AS04", "HP:0001945": "A03", "HP:0000458": "NS09"}
    drafts = {
        "HP:0000458": {"icpc3": "NS08", "icpc2": ""},
        "HP:0002027": {"icpc3": "DS01", "icpc2": ""},
    }
    table = {r["hpo_id"]: r for r in build_icpc_table.build(marc, drafts, BROWSER, str(obo))}
    assert table["HP:0012378"] == {
        "hpo_id": "HP:0012378", "hpo_term": "Fatigue", "icpc3": "AS04", "icpc2": "A04", "source": "marc",
    }
    assert (table["HP:0001945"]["icpc3"], table["HP:0001945"]["source"]) == ("AS03", "draft")
    assert (table["HP:0000458"]["icpc2"], table["HP:0000458"]["source"]) == ("N16", "draft")
    assert table["HP:0002027"]["source"] == "draft"


def test_build_refuses_a_code_with_several_icpc2_and_no_choice(tmp_path):
    obo = tmp_path / "hp.obo"
    obo.write_text(OBO)
    with pytest.raises(SystemExit, match="choose one"):
        build_icpc_table.build({"HP:0012378": "NS11"}, {}, BROWSER, str(obo))


def test_wheel_holds_the_table(tmp_path):
    subprocess.run(["uv", "build", "--wheel", "-o", str(tmp_path)], cwd=ROOT, capture_output=True, check=True)
    wheel = next(tmp_path.glob("*.whl"))
    assert "phenoscribe/resources/hpo_icpc.csv" in zipfile.ZipFile(wheel).namelist()


def test_build_refuses_an_icpc2_choice_outside_the_conversion(tmp_path):
    obo = tmp_path / "hp.obo"
    obo.write_text(OBO)
    drafts = {"HP:0012378": {"icpc3": "NS11", "icpc2": "A04"}}
    with pytest.raises(SystemExit, match="not a conversion"):
        build_icpc_table.build({}, drafts, BROWSER, str(obo))


def test_provenance_keeps_one_icpc_row(tmp_path):
    path = tmp_path / "out.xlsx"
    for _ in range(2):
        write_excel("MGA.001", [{"hpo_id": "HP:0012378", "hpo_term": "Fatigue", "patient_verbatim": ""}],
                    str(path), fmt="detailed", hpo_release="hp/releases/2026-02-16")
    keys = [r[0] for r in openpyxl.load_workbook(path)["Provenance"].iter_rows(values_only=True)]
    assert keys == ["Key", "hpo_release", "icpc_table"]


def test_add_icpc_columns_after_a_named_column_keeps_review_fill(tmp_path):
    from openpyxl.styles import PatternFill

    src = tmp_path / "in.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["HPO Code", "Mentions", "Code Source"])
    ws.append(["HP:0012378", 3, "hand"])
    for col in (1, 2, 3):
        ws.cell(row=2, column=col).fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    wb.save(src)

    add_icpc_columns.add_icpc_columns(str(src), str(src), after="Mentions")

    ws = openpyxl.load_workbook(src).active
    assert [c.value for c in ws[1]] == ["HPO Code", "Mentions", *ICPC_HEADERS, "Code Source"]
    assert [c.value for c in ws[2]] == ["HP:0012378", 3, *FATIGUE.values(), "hand"]
    assert ws.cell(row=2, column=3).fill.start_color.rgb.endswith("FFF2CC")


def test_add_icpc_columns_replaces_the_old_columns(tmp_path):
    src = tmp_path / "in.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["HPO Code", "Mentions", "ICPC-3", "ICPC-2", "ICPC Source", "Code Source"])
    ws.append(["HP:0012378", 3, "AS50", "A91", "draft", "hand"])
    wb.save(src)

    add_icpc_columns.add_icpc_columns(str(src), str(src), after="Mentions")

    rows = list(openpyxl.load_workbook(src).active.iter_rows(values_only=True))
    assert rows[0] == ("HPO Code", "Mentions", *ICPC_HEADERS, "Code Source")
    assert rows[1] == ("HP:0012378", 3, *FATIGUE.values(), "hand")
