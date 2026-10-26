"""Tests for the place of each finding in the interview (timeline.py)."""

from openpyxl import load_workbook

from phenoscribe import pipeline
from phenoscribe.config import Config
from phenoscribe.extract_symptoms import _parse_response
from phenoscribe.output import MENTIONS_HEADERS, MENTIONS_SHEET
from phenoscribe.review import Review, approve
from phenoscribe.timeline import Transcript, add_timeline, mention_rows, time_label

TEXT = (
    "[00:00] Bonjour, qu'est-ce qui vous amène ?\n"
    "[00:05] J'ai mal à la tête depuis trois semaines.\n"
    "[02:10] Je suis très fatiguée le soir.\n"
    "[05:30] Et la tête, ça me lance encore, j'ai mal à la tête le matin.\n"
    "[1:02:03] Pas de fièvre."
)

HEADACHE = {"hpo_id": "HP:0002315", "hpo_term": "Headache", "clinical_term": "headache",
            "negated": False, "patient_verbatim": "J'ai mal à la tête",
            "mentions": ["J'ai mal à la tête depuis trois semaines", "j'ai mal à la tête le matin"]}
FATIGUE = {"hpo_id": "HP:0012378", "hpo_term": "Fatigue", "clinical_term": "fatigue",
           "negated": False, "patient_verbatim": "très fatiguée",
           "mentions": ["je suis tres fatiguee le soir"]}  # accents dropped by the LLM
FEVER = {"hpo_id": "HP:0001945", "hpo_term": "Fever", "clinical_term": "fever",
         "negated": True, "patient_verbatim": "pas de fièvre", "mentions": ["Pas de fièvre."]}


def test_position_counts_spoken_text_only():
    with_marks = Transcript("[00:00] Bonjour. [00:10] Merci.")
    without = Transcript("Bonjour. Merci.")
    a = with_marks.locate_all("Merci")[0][0]
    b = without.locate_all("Merci")[0][0]
    assert with_marks.position(a) == without.position(b)
    assert with_marks.position(0) == 0.0 and with_marks.position(len(with_marks.text)) == 100.0


def test_time_comes_from_the_last_marker_before_the_quote():
    t = Transcript(TEXT)
    start = t.locate_all("Pas de fièvre")[0][0]
    assert t.seconds_at(start) == 3723
    assert time_label(3723) == "1:02:03" and time_label(130) == "02:10" and time_label(None) == ""
    assert Transcript("Pas de marqueurs.").seconds_at(5) is None


def test_exact_approximate_and_not_found():
    t = Transcript(TEXT)
    assert t.locate_all("je suis tres fatiguee le soir")[0][2] == "exact"  # accents fold
    assert t.locate_all("je suis vraiment très fatiguée le soir")[0][2] == "approximate"
    assert t.locate_all("douleur thoracique à l'effort") == []


def test_mention_rows_in_interview_order_with_numbers():
    rows = mention_rows([FEVER, HEADACHE, FATIGUE], TEXT)
    summary = [(r["hpo_term"], r["mention"], r["order"], r["seconds"]) for r in rows]
    assert summary == [
        ("Headache", 1, 1, 5),
        ("Fatigue", 1, 2, 130),
        ("Headache", 2, 1, 330),
        ("Fever", 1, 3, 3723),
    ]
    positions = [r["position"] for r in rows]
    assert positions == sorted(positions) and 0 < positions[0] < positions[-1] < 100


def test_a_short_quote_inside_a_long_one_is_one_mention():
    headache = dict(HEADACHE, mentions=["J'ai mal à la tête depuis trois semaines", "mal à la tête depuis"])
    rows = [r for r in mention_rows([headache], TEXT) if r["position"] is not None]
    # "mal à la tête depuis" overlaps the first quote. The short quote also
    # finds no second place, so one mention stays.
    assert len(rows) == 1


def test_a_repeated_quote_counts_each_place():
    text = "J'ai mal. Toujours mal. Encore mal."
    rows = mention_rows([{"hpo_id": "HP:0012531", "hpo_term": "Pain", "negated": False,
                          "mentions": ["mal"]}], text)
    assert [r["mention"] for r in rows] == [1, 2, 3]


def test_a_quote_not_found_keeps_an_empty_position_at_the_end():
    lost = dict(FATIGUE, mentions=["quelque chose que personne n'a dit"])
    rows = mention_rows([lost, HEADACHE], TEXT)
    assert rows[-1]["found"] == "not found" and rows[-1]["position"] is None
    assert rows[-1]["order"] is None


def test_no_mentions_falls_back_to_the_verbatim():
    rows = mention_rows([dict(FEVER, mentions=[])], TEXT)
    assert rows[0]["quote"] == "pas de fièvre" and rows[0]["found"] == "exact"


def test_add_timeline_puts_first_last_and_count_on_each_match():
    matches = [dict(HEADACHE), dict(FATIGUE)]
    add_timeline(matches, mention_rows(matches, TEXT))
    headache, fatigue = matches
    assert headache["n_mentions"] == 2 and headache["first_position"] < headache["last_position"]
    assert fatigue["n_mentions"] == 1 and fatigue["first_position"] == fatigue["last_position"]


def test_the_extractor_keeps_the_mention_list():
    found = _parse_response(
        '[{"patient_verbatim": "j\'ai mal", "clinical_term": "pain", "mentions": ["j\'ai mal", "encore mal"]},'
        ' {"patient_verbatim": "fatiguée", "clinical_term": "fatigue"},'
        ' {"patient_verbatim": "", "clinical_term": "cough", "mentions": "je tousse"}]'
    )
    assert [f["mentions"] for f in found] == [["j'ai mal", "encore mal"], ["fatiguée"], ["je tousse"]]


def test_the_excel_gets_a_mentions_sheet_and_three_columns(tmp_path, monkeypatch):
    monkeypatch.setattr(pipeline, "check_obo_version", lambda *a: "hp/releases/2026-02-16")
    monkeypatch.setattr(pipeline, "extract_symptoms", lambda text, **k: [{}])
    monkeypatch.setattr(pipeline, "match_hpo", lambda symptoms, **k: [dict(HEADACHE), dict(FEVER)])
    config = Config()
    config.output.format = "detailed"
    config.output.propagate_ancestors = False
    review = approve(Review.from_spans("467.txt", "MGA.467", TEXT, []))
    excel = tmp_path / "results.xlsx"
    pipeline.code_review(review, config, excel)

    wb = load_workbook(excel)
    main = wb.worksheets[0]
    header = [c.value for c in main[1]]
    first = dict(zip(header, [c.value for c in main[2]]))
    assert first["HPO Code"] == "HP:0002315" and first["Mentions"] == 2
    assert 0 < first["First Mention %"] < first["Last Mention %"] < 100

    sheet = wb[MENTIONS_SHEET]
    assert [c.value for c in sheet[1]] == MENTIONS_HEADERS
    rows = [dict(zip(MENTIONS_HEADERS, [c.value for c in row])) for row in sheet.iter_rows(min_row=2)]
    assert [(r["HPO Term"], r["Mention"], r["Time"]) for r in rows] == [
        ("Headache", 1, "00:05"), ("Headache", 2, "05:30"), ("Fever", 1, "1:02:03"),
    ]
    assert rows[-1]["Present/Absent"] == "Absent"
