"""Tests for the review of personal details and the gate before the LLM."""

import pytest
from docx import Document
from openpyxl import load_workbook

from phenoscribe import pipeline
from phenoscribe.config import Config, PatientConfig
from phenoscribe.output import PERSONAL_DETAILS_SHEET
from phenoscribe.pii import Span
from phenoscribe.review import (
    APPROVED,
    AUTO,
    BULK,
    DONE,
    NEEDS_REVIEW,
    Review,
    ReviewError,
    add_value,
    approve,
    clean_text,
    entries,
    fold,
    leaks,
    load_review,
    possibly_missed,
    review_path,
    save_review,
    set_keep_types,
    set_removed,
    summary_line,
    warnings,
)
from phenoscribe.transcribe import read_docx

TEXT = (
    "Bonjour madame Dupont. Mme DUPONT a 54 ans et elle est institutrice. "
    "Le docteur Vandermeulen la suit à Namur. Pas de mariage prévu, Marie dit non."
)


def _span(sub, category, score=0.95, origin="model", text=TEXT):
    start = text.index(sub)
    return Span(start, start + len(sub), sub, category, score, origin)


def _review(**kwargs):
    spans = [
        _span("Dupont", "PERSON", 0.98),
        _span("54 ans", "AGE", 0.8),
        _span("institutrice", "JOB", 0.7),
        _span("Namur", "LOCATION", 0.9),
        _span("Marie", "PERSON", 0.65),
    ]
    return Review.from_spans("412.txt", "MGA.412", TEXT, spans, **kwargs)


def test_fold_keeps_the_length():
    assert fold("Élodie DUPONT") == "elodie dupont"
    assert len(fold("Œuvre ﬁn İ")) == len("Œuvre ﬁn İ")


def test_a_value_found_once_is_removed_at_every_position_in_any_case():
    text, mapping = clean_text(_review())
    assert "Dupont" not in text and "DUPONT" not in text
    assert text.count("PERSON_1") == 2
    assert mapping["PERSON_1"] == "Dupont"


def test_whole_words_only():
    text, _ = clean_text(_review())
    assert "mariage" in text  # "Marie" does not hit "mariage"
    assert "PERSON_2 dit non" in text


def test_untick_puts_the_value_back_everywhere():
    review = set_removed(_review(), "Dupont", False)
    text, _ = clean_text(review)
    assert text.count("Dupont") == 1 and "DUPONT" in text
    assert [e.removed for e in entries(review) if e.value == "Dupont"] == [False]
    assert set_removed(review, "Dupont", True) and "Dupont" not in clean_text(review)[0]


def test_an_added_word_is_removed_and_can_be_put_back():
    review = add_value(_review(), "vandermeulen", "PERSON")
    assert "Vandermeulen" not in clean_text(review)[0]
    assert possibly_missed(review) == []
    review = set_removed(review, "vandermeulen", False)
    assert "Vandermeulen" in clean_text(review)[0]
    assert review.added == []


def test_kept_types_stay_and_names_still_go():
    review = _review(keep_types=["AGE", "JOB"])
    text, _ = clean_text(review)
    assert "54 ans" in text and "institutrice" in text and "Dupont" not in text
    assert leaks(review) == []
    # One value of a kept type can still be removed.
    text, _ = clean_text(set_removed(review, "institutrice", True))
    assert "institutrice" not in text and "54 ans" in text


def test_names_cannot_be_kept_as_a_type():
    review = set_keep_types(_review(), ["PERSON", "AGE"])
    assert review.keep_types == ["AGE"]
    assert "Dupont" not in clean_text(review)[0]


def test_overlapping_and_repeated_spans():
    text = "Jean Dupont et Dupont."
    spans = [Span(0, 11, "Jean Dupont", "PERSON"), Span(5, 11, "Dupont", "PERSON"),
             Span(15, 21, "Dupont", "PERSON")]
    clean, mapping = clean_text(Review.from_spans("a.txt", "A", text, spans))
    assert clean == "PERSON_1 et PERSON_2."
    assert mapping == {"PERSON_1": "Jean Dupont", "PERSON_2": "Dupont"}


def test_an_elision_is_not_the_value():
    text = "Parce qu'on m'avait dit que M. Dupont m'a vu. Le M est là."
    review = Review.from_spans("a.txt", "A", text, [Span(text.index("M est"), text.index("M est") + 1, "M", "PERSON")])
    clean, _ = clean_text(review)
    assert "m'avait" in clean and "m'a vu" in clean
    assert "Le PERSON_1 est" in clean


def test_regex_only_spans():
    text = "Appelez le 0472 12 34 56 ou le 0472 12 34 56."
    start = text.index("0472")
    span = Span(start, start + 13, "0472 12 34 56", "PHONE", 1.0, "regex")
    clean, _ = clean_text(Review.from_spans("a.txt", "A", text, [span]))
    assert clean == "Appelez le PHONE_1 ou le PHONE_1."


def test_possibly_missed_finds_words_after_a_title_only():
    missed = possibly_missed(_review())
    assert [word for word, _ in missed] == ["Vandermeulen"]
    assert "docteur Vandermeulen" in missed[0][1]


def test_a_title_at_the_end_of_a_sentence_is_not_a_name():
    text = "Il faut trouver un docteur. Oui. Puis M. Lefèvre et Dr Haddad."
    review = Review.from_spans("a.txt", "A", text, [])
    assert [word for word, _ in possibly_missed(review)] == ["Lefèvre", "Haddad"]


def test_summary_line():
    assert summary_line(_review()) == "5 removed: 2 names · 1 place · 1 age · 1 job title"


def test_approve_and_an_edit_after_it():
    review = approve(_review())
    assert review.status == APPROVED and review.approved_at
    add_value(review, "Vandermeulen")
    assert review.status == NEEDS_REVIEW and review.approved_at is None


def test_approve_is_refused_while_a_value_leaks(monkeypatch):
    from phenoscribe import review as review_module

    # The leak check is a safety net: force a clean text that still holds the name.
    monkeypatch.setattr(review_module, "clean_text", lambda r: (r.raw_text, {}))
    review = _review()
    assert "Dupont" in leaks(review)
    with pytest.raises(ReviewError, match="Still in the clean text"):
        approve(review)
    assert review.status == NEEDS_REVIEW


def test_a_value_inside_a_word_goes_at_its_detected_position():
    review = Review.from_spans("a.txt", "A", "abcDupontdef", [Span(3, 9, "Dupont", "PERSON")])
    assert clean_text(review)[0] == "abcPERSON_1def"


def test_save_load_and_a_damaged_file(tmp_path):
    review = _review()
    save_review(review, tmp_path)
    assert load_review(tmp_path, "MGA.412") == review
    review_path(tmp_path, "MGA.412").write_text("{not json", encoding="utf-8")
    with pytest.raises(ReviewError, match="damaged"):
        load_review(tmp_path, "MGA.412")


def test_long_transcript_stays_fast():
    import time

    text = (TEXT + " ") * 400  # about 60 000 characters, a 2-hour interview
    spans = [Span(i, i + 6, "Dupont", "PERSON") for i in range(len(text)) if text.startswith("Dupont", i)]
    review = Review.from_spans("a.txt", "A", text, spans)
    began = time.perf_counter()
    clean, _ = clean_text(review)
    leaks(review)
    assert "Dupont" not in clean
    assert time.perf_counter() - began < 5


# --- .docx input ---------------------------------------------------------------------


def test_read_docx_keeps_headers_paragraphs_and_tables_in_order(tmp_path):
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "Dossier de Mme Lambert"
    header_table = doc.sections[0].header.add_table(rows=1, cols=1, width=doc.sections[0].page_width)
    header_table.cell(0, 0).text = "Patient: Claire Lambert"
    doc.add_paragraph("Premier paragraphe.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Téléphone", "0472 12 34 56"
    doc.add_paragraph("Dernier paragraphe.")
    path = tmp_path / "chemin avec espaces é" / "412.docx"
    path.parent.mkdir()
    doc.save(str(path))
    assert read_docx(path) == (
        "Dossier de Mme Lambert\nPatient: Claire Lambert\nPremier paragraphe.\nTéléphone | 0472 12 34 56\nDernier paragraphe."
    )


def _plan(tmp_path, *names):
    paths = []
    for name in names:
        path = tmp_path / name
        path.write_bytes(b"x")
        paths.append(str(path))
    plans = pipeline.plan_batch(paths, Config(patient=PatientConfig(id_prefix="MGA.")), tmp_path)
    return {p.name: (p.action, p.note) for p in plans}


def test_plan_docx_is_a_transcript(tmp_path):
    assert _plan(tmp_path, "412.docx", "412.mp3", "~$412.docx", "old.doc") == {
        "412.docx": ("transcript", ""),
        "412.mp3": ("skip", "Uses 412.docx instead"),
        "old.doc": ("skip", "Old Word file. Save it as .docx in Word, then add it again"),
    }


def test_plan_txt_and_docx_for_one_patient_are_both_skipped(tmp_path):
    plans = _plan(tmp_path, "412.txt", "412.docx")
    assert {action for action, _ in plans.values()} == {"skip"}


# --- the gate before the LLM ---------------------------------------------------------


@pytest.fixture
def stub_steps(monkeypatch):
    """Stub the detector and both LLM steps; record the text the LLM gets."""
    sent = []
    monkeypatch.setattr(pipeline, "check_obo_version", lambda *a: "hp/releases/2026-02-16")
    monkeypatch.setattr(pipeline, "find_personal_details",
                        lambda text, **k: [_span("Dupont", "PERSON", text=text)] if "Dupont" in text else [])
    monkeypatch.setattr(pipeline, "extract_symptoms", lambda text, **k: sent.append(text) or [])
    monkeypatch.setattr(pipeline, "match_hpo", lambda symptoms, **k: [])
    return sent


def _transcript(tmp_path):
    path = tmp_path / "in" / "412.txt"
    path.parent.mkdir()
    path.write_text("Bonjour madame Dupont, je suis fatiguée.", encoding="utf-8")
    return str(path)


def test_no_llm_call_before_approval(tmp_path, stub_steps):
    out = str(tmp_path / "out" / "results.xlsx")
    with pytest.raises(pipeline.NeedsReview):
        pipeline.process_recording(_transcript(tmp_path), "MGA.412", Config(), output_path=out)
    assert stub_steps == []
    assert load_review(tmp_path / "out", "MGA.412").status == NEEDS_REVIEW


def test_the_llm_gets_only_the_approved_clean_text(tmp_path, stub_steps):
    out_dir = tmp_path / "out"
    source = _transcript(tmp_path)
    with pytest.raises(pipeline.NeedsReview):
        pipeline.process_recording(source, "MGA.412", Config(), output_path=str(out_dir / "r.xlsx"))
    review = load_review(out_dir, "MGA.412")
    add_value(review, "fatiguée", "PERSON")  # the person removes one more word
    save_review(approve(review), out_dir)

    pipeline.process_recording(source, "MGA.412", Config(), output_path=str(out_dir / "r.xlsx"))
    assert stub_steps == ["Bonjour madame PERSON_1, je suis PERSON_2."]
    assert load_review(out_dir, "MGA.412").status == DONE


def test_no_review_approves_and_records_it(tmp_path, stub_steps):
    out_dir = tmp_path / "out"
    pipeline.process_recording(_transcript(tmp_path), "MGA.412", Config(),
                                output_path=str(out_dir / "r.xlsx"), no_review=True)
    assert stub_steps == ["Bonjour madame PERSON_1, je suis fatiguée."]
    assert load_review(out_dir, "MGA.412").approved_how == "no review (CLI)"
    # A later run without the flag asks for a review of the unchecked file.
    with pytest.raises(pipeline.NeedsReview):
        pipeline.process_recording(str(tmp_path / "in" / "412.txt"), "MGA.412", Config(),
                                   output_path=str(out_dir / "r.xlsx"))


def test_code_batch_codes_only_approved_files_and_adds_the_counts_sheet(tmp_path, stub_steps):
    out_dir = tmp_path / "out"
    results = pipeline.prepare_batch([_transcript(tmp_path)], Config(), out_dir)
    assert results[0].review.status == NEEDS_REVIEW and stub_steps == []

    assert pipeline.code_batch(["MGA.412"], Config(), out_dir).files[0].error is not None
    assert stub_steps == []

    save_review(approve(load_review(out_dir, "412")), out_dir)
    batch = pipeline.code_batch(["412"], Config(), out_dir)
    assert stub_steps == ["Bonjour madame PERSON_1, je suis fatiguée."]
    sheet = load_workbook(batch.excel_path)[PERSONAL_DETAILS_SHEET]
    header = [c.value for c in sheet[1]]
    row = dict(zip(header, [c.value for c in sheet[2]]))
    assert row["Patient_ID"] == "412" and row["Names"] == 1 and row["Approved how"] == "checked"
    assert "Dupont" not in str(list(sheet.values))


def test_finish_clean(tmp_path, stub_steps):
    out_dir = tmp_path / "out"
    pipeline.prepare_batch([_transcript(tmp_path)], Config(), out_dir)
    assert pipeline.finish_clean(["412"], out_dir) is None  # nothing approved
    save_review(approve(load_review(out_dir, "412")), out_dir)
    archive = pipeline.finish_clean(["412"], out_dir)

    import zipfile

    names = zipfile.ZipFile(archive).namelist()
    assert "clean/412.txt" in names and any(n.startswith("personal_details_") for n in names)
    assert load_review(out_dir, "412").status == DONE
    assert (out_dir / "clean" / "412.txt").exists()


# --- doubtful files only ---------------------------------------------------------------


def test_warnings_name_the_reasons_to_check():
    assert warnings(_review()) == [
        "Possibly missed: Vandermeulen",
        "Unsure about: Marie, institutrice",
    ]
    nothing = Review.from_spans("a.txt", "A", "Je suis fatiguée.", [])
    assert warnings(nothing) == ["No personal detail found"]
    clear = Review.from_spans("a.txt", "A", "Madame Dupont va mieux.",
                              [Span(7, 13, "Dupont", "PERSON", 0.97)])
    assert warnings(clear) == []


def test_clear_files_are_approved_and_doubtful_files_wait(tmp_path):
    clear = Review.from_spans("a.txt", "A", "Madame Dupont va mieux.",
                              [Span(7, 13, "Dupont", "PERSON", 0.97)])
    doubtful = _review()
    assert pipeline.approve_clear_files([clear, doubtful], tmp_path) == 1
    assert clear.status == APPROVED and clear.approved_how == AUTO
    assert doubtful.status == NEEDS_REVIEW
    assert (tmp_path / "clean" / "A.txt").read_text(encoding="utf-8") == "Madame PERSON_1 va mieux."

    save_review(doubtful, tmp_path)
    assert pipeline.approve_the_rest(["A", "MGA.412"], tmp_path) == 1
    assert load_review(tmp_path, "MGA.412").approved_how == BULK


def test_an_old_review_file_still_loads(tmp_path):
    import json

    # The first version of the file: a "reviewed" flag, no "approved_how".
    save_review(approve(_review()), tmp_path)
    path = review_path(tmp_path, "MGA.412")
    data = json.loads(path.read_text(encoding="utf-8"))
    del data["approved_how"]
    data["reviewed"] = False
    path.write_text(json.dumps(data), encoding="utf-8")
    old = load_review(tmp_path, "MGA.412")
    assert old.status == APPROVED and old.approved_how == "" and old.kept == []
    assert "Dupont" not in clean_text(old)[0]


def test_a_review_file_with_bad_bytes_is_damaged(tmp_path):
    save_review(_review(), tmp_path)
    review_path(tmp_path, "MGA.412").write_bytes(b"\xff\xfe{")
    with pytest.raises(ReviewError, match="damaged"):
        load_review(tmp_path, "MGA.412")


def test_separators_and_unicode_case_do_not_hide_a_value():
    text = "Jean Dupont vient. Jean-Dupont repart. ΟΣ et ος."
    spans = [Span(0, 11, "Jean Dupont", "PERSON"), Span(39, 41, "ΟΣ", "PERSON")]
    review = Review.from_spans("a.txt", "A", text, spans)
    clean, _ = clean_text(review)
    assert clean == "PERSON_1 vient. PERSON_1 repart. PERSON_2 et PERSON_2."
    assert leaks(review) == []


def test_cli_clean_keeps_a_name_in_the_file_name_out_of_the_zip(tmp_path, stub_steps):
    import argparse
    import zipfile

    from phenoscribe import cli

    folder = tmp_path / "in"
    folder.mkdir()
    (folder / "martin-jean.txt").write_text("Bonjour madame Dupont.", encoding="utf-8")
    cli._cmd_clean(argparse.Namespace(input_dir=str(folder), output=str(tmp_path / "out"),
                                      config=str(tmp_path / "none.yaml")))
    (archive,) = (tmp_path / "out").glob("clean_texts_*.zip")
    names = zipfile.ZipFile(archive).namelist()
    assert not any("martin" in n for n in names)
    assert any(n.startswith("clean/pt-") for n in names)
