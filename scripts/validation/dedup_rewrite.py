#!/usr/bin/env python
"""Rewrite results.xlsx applying the merged dedup fix, no pipeline re-run.

The shipped workbook was written before the dedup fix merged, so it still has
duplicate (patient, HPO code) present rows that pollute the closure. This script
reconstructs the present matches (their real code is intact in the workbook),
runs them through the merged `dedup_by_code`, preserves absent rows 1:1, and
rebuilds the Ancestor Closure from the deduped present set — all via the merged
output code path. Clinical Term stays blank (it was never stored; cannot backfill).
"""
import argparse
import shutil
import sys
from collections import OrderedDict, defaultdict
from pathlib import Path

import openpyxl
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "src"))
from phenoscribe import output as O  # noqa: E402

ABSENT_CODE = O._ABSENT_CODE_PLACEHOLDER


def row_to_match(h, r, negated):
    g = lambda name: r[h[name]] if name in h else None
    conf = g("Confidence")
    return {
        "hpo_term": g("HPO Term") or "",
        "hpo_id": None if negated else (str(g("HPO Code")).strip() if g("HPO Code") else None),
        "negated": negated,
        "frequency": g("Frequency") or "",
        "onset": g("Onset") or "",
        "severity": g("Severity") or "",
        "patient_verbatim": g("Patient Verbatim") or "",
        "clinical_term": (g("Clinical Term") or "") if "Clinical Term" in h else "",
        "needs_review": str(g("Needs Review") or "").strip().upper() == "REVIEW",
        "confidence": float(conf) if isinstance(conf, (int, float)) else None,
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--results", default="output/results.xlsx")
    ap.add_argument("--obo", default="data/hpo/hp.obo")
    ap.add_argument("--release", default="hp/releases/2026-02-16")
    ap.add_argument("--fmt", default="detailed")
    args = ap.parse_args()

    src = Path(args.results)
    backup = src.with_suffix(".pre-dedup.xlsx")
    shutil.copy2(src, backup)
    print(f"backed up {src} -> {backup}")

    wb_old = openpyxl.load_workbook(src, data_only=True)
    rows = list(wb_old["Sheet"].iter_rows(values_only=True))
    h = {name: i for i, name in enumerate(rows[0])}

    present = defaultdict(list)
    absent = defaultdict(list)
    order = OrderedDict()  # preserve patient order
    for r in rows[1:]:
        if r[h["Patient_ID"]] is None:
            continue
        pid = str(r[h["Patient_ID"]]).strip()
        order[pid] = True
        is_absent = str(r[h["Present/Absent"]] or "").strip().lower() == "absent"
        (absent if is_absent else present)[pid].append(row_to_match(h, r, is_absent))

    # Fresh workbook via the merged writer helpers.
    wb = Workbook()
    ws = wb.active
    headers = O.HEADERS[args.fmt]
    ws.append(headers)
    O._style_header_row(ws, len(headers))

    deduped_present = {}
    before = after = 0
    for pid in order:
        dp = O.dedup_by_code(present[pid], args.fmt)
        deduped_present[pid] = dp
        before += len(present[pid])
        after += len(dp)
        O._write_detailed_format(ws, pid, dp)
        O._write_detailed_format(ws, pid, absent[pid])  # absent preserved 1:1

    O._auto_fit_columns(ws)
    O._stamp_provenance(wb, args.release)
    for pid in order:
        O._write_closure_sheet(wb, pid, deduped_present[pid], args.obo)

    wb.save(src)
    print(f"present rows: {before} -> {after} (removed {before - after} duplicate restatements)")
    print(f"absent rows preserved: {sum(len(v) for v in absent.values())}")
    print(f"rewrote {src}")


if __name__ == "__main__":
    main()
