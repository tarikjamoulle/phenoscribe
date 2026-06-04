#!/usr/bin/env python
"""Reference-free validation battery (source A) + raw Marc concordance counts (source B).

Read-only. Computes objectively-checkable properties of a Phenoscribe results
workbook that need NO gold standard, plus the raw overlap with Marc's manual
codes (reported as agreement, never accuracy — see the 2026-06-04 plan).

Usage:
  python scripts/validation/measure_reffree.py \
      --results output/results.xlsx \
      --gt data/ground_truth/CRS__under20_35_HPO_corr.xlsx \
      --obo data/hpo/hp.obo \
      --out output/validation/reffree.json \
      --disagreements output/validation/marc_disagreements.json

Exit code is always 0 (measurement, not a gate). Findings are in the JSON.
"""
import argparse
import json
import re
import statistics
import sys
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

# Reuse the pipeline's own HPO + IC machinery so we score with the same graph.
sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "src"))
from phenoscribe import validate as V  # noqa: E402
from phenoscribe.semantic_similarity import get_ic_map  # noqa: E402

# Marc writes some codes with the 'HPO:' prefix; accept both, normalise to HP:#######.
HP_RE = re.compile(r"HPO?:(\d{7})")


def _normalise_codes(text):
    return {"HP:" + m.group(1) for m in HP_RE.finditer(text)}


# 'Needs Review' carries string flags ('REVIEW', 'Yes', 'true', '1'), not bools.
_TRUTHY_FLAGS = {"true", "yes", "y", "1", "review", "needs review", "x"}


def _is_truthy_flag(v):
    if v is True:
        return True
    if v is False or v is None:
        return False
    if isinstance(v, (int, float)):
        return v != 0
    return str(v).strip().lower() in _TRUTHY_FLAGS


def _as_float(v):
    """Coerce numeric cells AND numeric strings ('0.83') to float; else None."""
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).strip())
    except (TypeError, ValueError):
        return None


def load_results(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    main = list(wb["Sheet"].iter_rows(values_only=True))
    hdr = main[0]
    idx = {name: i for i, name in enumerate(hdr)}
    present = defaultdict(list)   # patient -> [code, ...] (present only, dups kept)
    absent = defaultdict(list)    # patient -> [code, ...] (absent/negated)
    needs_review = 0
    confidences = []
    rows_total = 0
    for r in main[1:]:
        if r[idx["Patient_ID"]] is None:
            continue
        rows_total += 1
        pid = str(r[idx["Patient_ID"]]).strip()
        code = r[idx["HPO Code"]]
        pa = (r[idx["Present/Absent"]] or "").strip().lower()
        if "Needs Review" in idx and _is_truthy_flag(r[idx["Needs Review"]]):
            needs_review += 1
        if "Confidence" in idx:
            conf = _as_float(r[idx["Confidence"]])
            if conf is not None:
                confidences.append(conf)
        if not code:
            continue
        code = str(code).strip()
        if pa == "present":
            present[pid].append(code)
        elif pa == "absent":
            absent[pid].append(code)
    closure = None
    if "Ancestor Closure" in wb.sheetnames:
        closure = defaultdict(lambda: {"leaf": [], "ancestor": []})
        for r in list(wb["Ancestor Closure"].iter_rows(values_only=True))[1:]:
            if r[0] is None:
                continue
            closure[str(r[0]).strip()][(r[3] or "").strip()].append(str(r[2]).strip())
    provenance = dict(list(wb["Provenance"].iter_rows(values_only=True))[1:]) if "Provenance" in wb.sheetnames else {}
    return {
        "present": present, "absent": absent, "closure": closure,
        "needs_review": needs_review, "confidences": confidences,
        "rows_total": rows_total, "provenance": provenance,
    }


def load_gt(path):
    gt = openpyxl.load_workbook(path, data_only=True).active
    out = {}
    for r in gt.iter_rows(values_only=True):
        if r[0] is None:
            continue
        pid = str(r[0]).strip()
        joined = " ".join(str(c) for c in r if c is not None)
        codes = _normalise_codes(joined)
        if codes:
            out[pid] = codes
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--results", default="output/results.xlsx")
    ap.add_argument("--gt", default="data/ground_truth/CRS__under20_35_HPO_corr.xlsx")
    ap.add_argument("--obo", default="data/hpo/hp.obo")
    ap.add_argument("--out", default="output/validation/reffree.json")
    ap.add_argument("--disagreements", default="output/validation/marc_disagreements.json")
    args = ap.parse_args()

    R = load_results(args.results)
    hpo = V._get_hpo()
    obsolete = V._get_obsolete_map(args.obo)
    ic_map = get_ic_map(hpo)

    def is_valid(code):
        return hpo.graph_aware.__class__ and code in ic_map or _term_exists(hpo, code)

    findings = {}

    # --- A1: code validity + obsolescence -------------------------------------
    all_present = [(p, c) for p, cs in R["present"].items() for c in cs]
    invalid, obs = [], []
    for p, c in all_present:
        if not _term_exists(hpo, c):
            invalid.append({"patient": p, "code": c})
        if c in obsolete:
            obs.append({"patient": p, "code": c, "replaced_by": obsolete[c]})
    findings["A1_code_validity"] = {
        "present_codes_total": len(all_present),
        "distinct_codes": len({c for _, c in all_present}),
        "invalid_codes": invalid,
        "obsolete_codes": obs,
        "pass": not invalid and not obs,
    }

    # --- A2: dedup (no duplicate (patient, present code) rows) -----------------
    dup = []
    for p, cs in R["present"].items():
        for code, n in Counter(cs).items():
            if n > 1:
                dup.append({"patient": p, "code": code, "count": n})
    findings["A2_dedup"] = {
        "duplicate_rows": sorted(dup, key=lambda d: -d["count"]),
        "n_duplicated_pairs": len(dup),
        "pass": not dup,
    }

    # --- A3: ancestor-closure correctness -------------------------------------
    a3 = {"checked_patients": 0, "patients_mismatched": [], "negated_leaked_into_closure": [], "pass": True}
    if R["closure"] is not None:
        for p, cs in R["present"].items():
            leaves = {c for c in cs if _term_exists(hpo, c)}
            expected_anc = set()
            for leaf in leaves:
                expected_anc |= _ancestors(hpo, leaf)
            expected_closure = leaves | expected_anc
            got = set(R["closure"][p]["leaf"]) | set(R["closure"][p]["ancestor"])
            a3["checked_patients"] += 1
            missing = expected_closure - got
            extra = got - expected_closure
            if missing or extra:
                a3["patients_mismatched"].append({
                    "patient": p, "n_missing": len(missing), "n_extra": len(extra),
                    "missing_sample": sorted(missing)[:8], "extra_sample": sorted(extra)[:8],
                })
            # negated leaves must NOT appear as a closure leaf/ancestor source
            neg_leaked = set(R["absent"].get(p, [])) & got
            # only flag negated codes that are NOT also independently present
            neg_leaked -= leaves
            if neg_leaked:
                a3["negated_leaked_into_closure"].append({"patient": p, "codes": sorted(neg_leaked)})
        a3["pass"] = not a3["patients_mismatched"] and not a3["negated_leaked_into_closure"]
    else:
        a3 = {"present": False, "note": "no Ancestor Closure sheet"}
    findings["A3_ancestor_closure"] = a3

    # --- A4: judge hygiene -----------------------------------------------------
    confs = R["confidences"]
    findings["A4_judge_hygiene"] = {
        "needs_review_rows": R["needs_review"],
        "rows_total": R["rows_total"],
        "confidence_median": round(statistics.median(confs), 3) if confs else None,
        "confidence_min": round(min(confs), 3) if confs else None,
    }

    # --- A5: negation surfaced -------------------------------------------------
    findings["A5_negation_surfaced"] = {
        "present_rows": sum(len(v) for v in R["present"].values()),
        "absent_rows": sum(len(v) for v in R["absent"].values()),
        "patients_with_absent": sorted(p for p, v in R["absent"].items() if v),
    }

    # --- A6: IC distribution of predicted (present) terms ----------------------
    ics = [ic_map.get(c) for _, c in all_present if c in ic_map]
    ics = [x for x in ics if x is not None]
    if ics:
        # low-IC = near-root "syntactic" predictions; threshold at IC<2 (~top of tree)
        low = [x for x in ics if x < 2.0]
        findings["A6_ic_distribution"] = {
            "n_scored": len(ics),
            "median": round(statistics.median(ics), 3),
            "mean": round(statistics.mean(ics), 3),
            "min": round(min(ics), 3),
            "max": round(max(ics), 3),
            "frac_low_ic_below_2": round(len(low) / len(ics), 3),
        }
    else:
        findings["A6_ic_distribution"] = {"n_scored": 0}

    # --- B (raw): concordance with Marc, the GT-matched patients ---------------
    GT = load_gt(args.gt)
    matched = sorted(set(R["present"]) & set(GT))
    conc = {"matched_patients": matched, "per_patient": {}, "micro": {}}
    tp = fp = fn = 0
    disagreements = {}
    for p in matched:
        pred = set(R["present"][p])
        gold = GT[p]
        both = pred & gold
        po = pred - gold      # phenoscribe-only
        mo = gold - pred      # marc-only
        tp += len(both); fp += len(po); fn += len(mo)
        conc["per_patient"][p] = {
            "pred_n": len(pred), "gold_n": len(gold),
            "agree_n": len(both), "pheno_only_n": len(po), "marc_only_n": len(mo),
            "agreement_jaccard": round(len(both) / len(pred | gold), 3) if (pred | gold) else None,
        }
        disagreements[p] = {
            "agree": sorted(both),
            "pheno_only": sorted(po),
            "marc_only": sorted(mo),
        }
    if tp + fp and tp + fn:
        prec = tp / (tp + fp); rec = tp / (tp + fn)
        conc["micro"] = {
            "label": "AGREEMENT with one annotator, NOT accuracy",
            "tp": tp, "fp": fp, "fn": fn,
            "precision_vs_marc": round(prec, 3),
            "recall_vs_marc": round(rec, 3),
            "f1_vs_marc": round(2 * prec * rec / (prec + rec), 3) if (prec + rec) else 0,
            "n_patients": len(matched),
        }
    findings["B_marc_concordance"] = conc
    findings["_meta"] = {
        "results": args.results, "provenance": R["provenance"],
        "note": "A* = reference-free (objective). B = agreement with Marc, never accuracy.",
    }

    Path(args.out).parent.mkdir(parents=True, exist_ok=True)
    Path(args.out).write_text(json.dumps(findings, indent=2, ensure_ascii=False))
    Path(args.disagreements).parent.mkdir(parents=True, exist_ok=True)
    Path(args.disagreements).write_text(json.dumps(disagreements, indent=2, ensure_ascii=False))

    # human summary to stdout
    print(f"results: {args.results}  release: {R['provenance'].get('hpo_release')}")
    print(f"A1 validity: {'PASS' if findings['A1_code_validity']['pass'] else 'FAIL'} "
          f"({len(invalid)} invalid, {len(obs)} obsolete of {len(all_present)} present codes)")
    print(f"A2 dedup:    {'PASS' if findings['A2_dedup']['pass'] else 'FAIL'} "
          f"({len(dup)} duplicated (patient,code) pairs)")
    print(f"A3 closure:  {'PASS' if findings['A3_ancestor_closure'].get('pass') else 'FAIL'} "
          f"({len(findings['A3_ancestor_closure'].get('patients_mismatched', []))} mismatched, "
          f"{len(findings['A3_ancestor_closure'].get('negated_leaked_into_closure', []))} negation-leaks)")
    print(f"A4 hygiene:  {R['needs_review']} needs-review of {R['rows_total']} rows")
    print(f"A5 negation: {findings['A5_negation_surfaced']['present_rows']} present / "
          f"{findings['A5_negation_surfaced']['absent_rows']} absent")
    a6 = findings["A6_ic_distribution"]
    if a6.get("n_scored"):
        print(f"A6 IC:       median {a6['median']}, {int(a6['frac_low_ic_below_2']*100)}% below IC 2.0 (near-root)")
    if conc["micro"]:
        m = conc["micro"]
        print(f"B  vs Marc:  AGREEMENT n={m['n_patients']}  P={m['precision_vs_marc']} "
              f"R={m['recall_vs_marc']} F1={m['f1_vs_marc']} (NOT accuracy)")
    print(f"\nwrote {args.out} and {args.disagreements}")


_TERM_CACHE = {}


def _term_exists(hpo, code):
    if code in _TERM_CACHE:
        return _TERM_CACHE[code]
    try:
        ok = hpo.get_term(code) is not None
    except Exception:
        ok = False
    _TERM_CACHE[code] = ok
    return ok


def _ancestors(hpo, code):
    try:
        return {a.value if hasattr(a, "value") else str(a)
                for a in hpo.graph.get_ancestors(code)}
    except Exception:
        try:
            return {str(a) for a in hpo.graph.get_ancestors(hpo.get_term(code).identifier)}
        except Exception:
            return set()


if __name__ == "__main__":
    main()
