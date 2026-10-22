"""Smoke run for CI: one short French text through the whole pipeline, LLM stubbed.

Run after `phenoscribe setup --no-models`. It proves that the locked packages
install and work together on this OS: PII model, ChromaDB search, Excel output.
The PII model downloads here because setup skipped it.

    uv run python scripts/smoke_run.py
"""

import json
import re
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

from phenoscribe import extract_symptoms, match_hpo
from phenoscribe.config import load_config
from phenoscribe.pipeline import process_recording

SAMPLE = Path(__file__).resolve().parent.parent / "tests" / "fixtures" / "sample.txt"

FINDINGS = [
    {"clinical_term": "Headache", "patient_verbatim": "maux de tête persistants"},
    {"clinical_term": "Fatigue", "patient_verbatim": "une fatigue importante"},
]


def fake_llm_call(system_prompt, user_prompt, **kwargs):
    """Extraction returns FINDINGS. The judge picks the first candidate."""
    if "Candidate HPO terms" not in user_prompt:
        return json.dumps(FINDINGS)
    name, hpo_id = re.search(r"1\. (.+) \((HP:\d+)\)", user_prompt).groups()
    return json.dumps({"hpo_id": hpo_id, "hpo_term": name, "match": True, "confidence": 0.9})


def main() -> int:
    extract_symptoms.llm_call = fake_llm_call
    match_hpo.llm_call = fake_llm_call

    config = load_config()
    excel = Path(tempfile.mkdtemp()) / "smoke.xlsx"
    matches = process_recording(str(SAMPLE), "smoke", config, output_path=str(excel))

    rows = list(load_workbook(excel).worksheets[0].iter_rows(min_row=2, values_only=True))
    codes = [row[2] for row in rows if row[2]]
    print(f"Smoke run: {len(matches)} matches, codes {codes}, Excel {excel}")
    return 0 if len(codes) == len(FINDINGS) else 1


if __name__ == "__main__":
    sys.exit(main())
